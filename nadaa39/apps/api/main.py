from __future__ import annotations

import base64
import io
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus, urlparse

import httpx
import numpy as np
import soundfile as sf
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document
from pptx import Presentation

load_dotenv()

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
DATA.mkdir(parents=True, exist_ok=True)
VOICE_FILE = DATA / "creator_embedding.npy"
UPLOADS = DATA / "uploads"
OUT = DATA / "out"
UPLOADS.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

app = FastAPI(title="Nadaa 39 MPI Gateway", version="0.1.0")
origins = [x.strip() for x in os.getenv("CORS_ORIGINS", "http://localhost:5173").split(",") if x.strip()]
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

SYSTEM = """You are Nadaa, pronounced nah-dah. You are the personal AI assistant created by Ramatema Pule.
You operate as the Morbius Processing Intelligence (MPI) orchestration layer.
Speak naturally, warmly and directly, like a highly capable human assistant, without pretending to be human.
When an action could have real-world consequences, explain the action and request confirmation before execution unless the platform explicitly marks it SAFE.
Never claim an IQ score as a scientific fact. '720' is Nadaa branding.
Prefer concise answers, but be thorough for complex work.
"""


class ChatBody(BaseModel):
    message: str
    context: dict[str, Any] = Field(default_factory=dict)


class ActionBody(BaseModel):
    action: str
    args: dict[str, Any] = Field(default_factory=dict)
    confirmed: bool = False


class LiveTokenBody(BaseModel):
    voice_name: str = "Aoede"


class VerifyBody(BaseModel):
    audio_base64: str


def _cosine(a: np.ndarray, b: np.ndarray) -> float:
    den = float(np.linalg.norm(a) * np.linalg.norm(b))
    return float(np.dot(a, b) / den) if den else 0.0


def _voice_model():
    try:
        from speechbrain.inference.speaker import EncoderClassifier
    except Exception as exc:
        raise RuntimeError("SpeechBrain is unavailable. Install speechbrain and matching torch/torchaudio.") from exc
    return EncoderClassifier.from_hparams(source="speechbrain/spkrec-ecapa-voxceleb")


def _embedding_from_bytes(audio_bytes: bytes) -> np.ndarray:
    model = _voice_model()
    with io.BytesIO(audio_bytes) as bio:
        audio, sr = sf.read(bio, dtype="float32", always_2d=False)
    if audio.ndim > 1:
        audio = audio.mean(axis=1)
    if sr != 16000:
        # Use librosa only when present; otherwise fail explicitly rather than silently mis-processing.
        try:
            import librosa
            audio = librosa.resample(audio, orig_sr=sr, target_sr=16000)
        except Exception as exc:
            raise RuntimeError("Audio must be 16 kHz mono, or librosa must be installed for resampling.") from exc
    import torch
    signal = torch.tensor(audio).unsqueeze(0)
    emb = model.encode_batch(signal).squeeze().detach().cpu().numpy().astype(np.float32)
    emb /= max(float(np.linalg.norm(emb)), 1e-8)
    return emb


def _open_url(url: str) -> dict[str, Any]:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise HTTPException(400, "Only http/https URLs are accepted.")
    try:
        r = httpx.get(url, timeout=20.0, follow_redirects=True, headers={"User-Agent": "Nadaa39/0.1"})
        r.raise_for_status()
    except Exception as exc:
        raise HTTPException(502, f"Could not fetch URL: {exc}") from exc
    soup = BeautifulSoup(r.text, "lxml")
    for node in soup(["script", "style", "noscript"]):
        node.decompose()
    text = " ".join(soup.stripped_strings)
    return {"url": str(r.url), "title": soup.title.get_text(" ", strip=True) if soup.title else "", "text": text[:50000]}


async def _provider_gemini(message: str, context: dict[str, Any]) -> str:
    from google import genai
    from google.genai import types
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY not configured")
    client = genai.Client(api_key=api_key)
    contents: list[Any] = [message]
    if context:
        contents.append(json.dumps(context, ensure_ascii=False))
    response = client.models.generate_content(model=os.getenv("GEMINI_MODEL", "gemini-3.5-flash"), contents=contents, config=types.GenerateContentConfig(system_instruction=SYSTEM))
    return response.text or "I did not receive a text response."


async def _provider_openai_compatible(base_url: str, key: str, model: str, message: str, context: dict[str, Any]) -> str:
    payload = {
        "model": model,
        "messages": [{"role": "system", "content": SYSTEM}, {"role": "user", "content": message + ("\n\nContext:\n" + json.dumps(context) if context else "")}],
        "temperature": 0.4,
    }
    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    r = await httpx.AsyncClient(timeout=40).post(base_url, headers=headers, json=payload)
    r.raise_for_status()
    data = r.json()
    return data["choices"][0]["message"]["content"]


async def mpi(message: str, context: dict[str, Any]) -> dict[str, Any]:
    """MPI = model router + policy + fallback. It does not literally merge model weights."""
    providers: list[tuple[str, Any]] = [("gemini", _provider_gemini)]
    if os.getenv("GROQ_API_KEY"):
        providers.append(("groq", lambda m, c: _provider_openai_compatible("https://api.groq.com/openai/v1/chat/completions", os.environ["GROQ_API_KEY"], "openai/gpt-oss-120b", m, c)))
    if os.getenv("OPENROUTER_API_KEY"):
        providers.append(("openrouter", lambda m, c: _provider_openai_compatible("https://openrouter.ai/api/v1/chat/completions", os.environ["OPENROUTER_API_KEY"], "openrouter/free", m, c)))
    errors = []
    for name, fn in providers:
        try:
            return {"provider": name, "text": await fn(message, context)}
        except Exception as exc:
            errors.append(f"{name}: {exc}")
    return {"provider": "offline", "text": "Nadaa is running in offline-safe mode. Add at least one model API key to enable reasoning.", "errors": errors}


def _make_pdf(text: str, filename: str) -> Path:
    path = OUT / filename
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Helvetica", 10)
    for raw in text.splitlines() or [""]:
        words, line = raw.split(), ""
        for word in words:
            trial = (line + " " + word).strip()
            if c.stringWidth(trial, "Helvetica", 10) > width - 80:
                c.drawString(40, y, line)
                y -= 14
                if y < 50:
                    c.showPage(); c.setFont("Helvetica", 10); y = height - 50
                line = word
            else:
                line = trial
        c.drawString(40, y, line); y -= 14
        if y < 50:
            c.showPage(); c.setFont("Helvetica", 10); y = height - 50
    c.save()
    return path


def _make_docx(text: str, filename: str) -> Path:
    path = OUT / filename
    doc = Document()
    for line in text.splitlines() or [""]:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def _make_pptx(text: str, filename: str) -> Path:
    path = OUT / filename
    prs = Presentation()
    chunks = [x.strip() for x in re.split(r"\n\s*\n", text) if x.strip()] or [text]
    for i, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Nadaa 39 — {i+1}"
        slide.placeholders[1].text = chunk[:4000]
    prs.save(path)
    return path


def _make_tex(text: str, filename: str) -> Path:
    path = OUT / filename
    escaped = text.replace("\\", r"\\textbackslash{}").replace("&", r"\\&").replace("%", r"\\%").replace("#", r"\\#").replace("_", r"\\_")
    path.write_text("\\documentclass{article}\n\\usepackage[utf8]{inputenc}\n\\begin{document}\n" + escaped.replace("\n", "\n\n") + "\n\\end{document}\n", encoding="utf-8")
    return path


@app.get("/health")
async def health():
    return {"name": "Nadaa", "pronunciation": "nah-dah", "mpi": True, "creator": os.getenv("CREATOR_NAME", "Ramatema Pule")}


def _detect_action(message: str) -> dict[str, Any] | None:
    text = message.strip()
    low = text.lower()
    if "open whatsapp" in low:
        return {"name": "open_whatsapp", "risk": "SAFE", "args": {}}
    m = re.search(r"\bcall(?: me)?(?: on| at)?\s+([+0-9][+0-9() .-]{6,})", text, re.I)
    if m:
        return {"name": "call_phone", "risk": "CONFIRM", "args": {"number": m.group(1).strip()}}
    m = re.search(r"\bcall\s+([A-Za-z][A-Za-z .'-]{1,50})$", text, re.I)
    if m:
        return {"name": "call_contact", "risk": "CONFIRM", "args": {"name": m.group(1).strip()}}
    if low.startswith(("create a pdf", "make a pdf", "create pdf")):
        return {"name":"generate_pdf","risk":"SAFE","args":{"text":text}}
    if low.startswith(("create a word", "create a docx", "make a word")):
        return {"name":"generate_docx","risk":"SAFE","args":{"text":text}}
    if low.startswith(("create a powerpoint", "create a ppt", "make a powerpoint")):
        return {"name":"generate_pptx","risk":"SAFE","args":{"text":text}}
    return None


@app.post("/api/chat")
async def chat(body: ChatBody):
    result = await mpi(body.message, body.context)
    action_request = _detect_action(body.message)
    if action_request:
        result["action"] = action_request
    return result


@app.post("/api/live-token")
async def live_token(body: LiveTokenBody):
    key = os.getenv("GEMINI_API_KEY")
    if not key:
        raise HTTPException(503, "GEMINI_API_KEY is not configured on the server.")
    try:
        from google import genai
        from datetime import datetime, timedelta, timezone
        client = genai.Client(api_key=key)
        expire = datetime.now(timezone.utc) + timedelta(minutes=30)
        token = client.auth_tokens.create(config={"uses": 1, "expire_time": expire, "live_connect_constraints": {"model": os.getenv("GEMINI_LIVE_MODEL", "gemini-3.1-flash-live-preview"), "config": {"response_modalities": ["AUDIO"], "system_instruction": {"parts": [{"text": SYSTEM}],}, "speech_config": {"voice_config": {"prebuilt_voice_config": {"voice_name": body.voice_name}}}}}})
        return {"token": token.name, "model": os.getenv("GEMINI_LIVE_MODEL", "gemini-3.1-flash-live-preview")}
    except Exception as exc:
        raise HTTPException(500, str(exc)) from exc


@app.post("/api/voice/enroll")
async def voice_enroll(file: UploadFile = File(...)):
    content = await file.read()
    try:
        emb = _embedding_from_bytes(content)
    except Exception as exc:
        raise HTTPException(400, str(exc)) from exc
    np.save(VOICE_FILE, emb)
    return {"ok": True, "creator": os.getenv("CREATOR_NAME", "Ramatema Pule"), "embedding_saved": True}


@app.post("/api/voice/verify")
async def voice_verify(body: VerifyBody):
    if not VOICE_FILE.exists():
        raise HTTPException(409, "Enroll the creator voice first.")
    try:
        raw = base64.b64decode(body.audio_base64)
        emb = _embedding_from_bytes(raw)
    except Exception as exc:
        raise HTTPException(400, str(exc)) from exc
    ref = np.load(VOICE_FILE)
    score = _cosine(ref, emb)
    threshold = float(os.getenv("VOICE_THRESHOLD", "0.72"))
    return {"creator_verified": score >= threshold, "score": score, "threshold": threshold}


@app.post("/api/fetch-url")
async def fetch_url(body: dict[str, str]):
    return _open_url(body.get("url", ""))


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    safe_name = Path(file.filename or "upload.bin").name
    path = UPLOADS / safe_name
    with path.open("wb") as handle:
        shutil.copyfileobj(file.file, handle)
    return {"filename": safe_name, "path": str(path), "content_type": file.content_type}


@app.post("/api/action")
async def action(body: ActionBody):
    destructive = {"call_phone", "send_message", "delete_file", "send_whatsapp"}
    if body.action in destructive and not body.confirmed:
        return {"status": "confirmation_required", "action": body.action, "args": body.args}
    if body.action == "open_url":
        url = body.args.get("url", "")
        if not url.startswith(("http://", "https://", "whatsapp://", "tel:")):
            raise HTTPException(400, "Unsupported URL scheme")
        return {"status": "handoff", "url": url}
    # Web clients cannot safely execute arbitrary desktop commands. Native shells should implement these actions.
    if body.action in destructive:
        return {"status": "native_handoff", "action": body.action, "args": body.args}
    return {"status": "unknown_action", "action": body.action}


@app.post("/api/generate/{kind}")
async def generate(kind: str, body: dict[str, str]):
    text = body.get("text", "")
    if not text.strip():
        raise HTTPException(400, "text is required")
    mapping = {"pdf": (_make_pdf, "nadaa-document.pdf"), "docx": (_make_docx, "nadaa-document.docx"), "pptx": (_make_pptx, "nadaa-presentation.pptx"), "tex": (_make_tex, "nadaa-document.tex")}
    if kind not in mapping:
        raise HTTPException(404, "Supported kinds: pdf, docx, pptx, tex")
    fn, name = mapping[kind]
    path = fn(text, name)
    return FileResponse(path, filename=name)
